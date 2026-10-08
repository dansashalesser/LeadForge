"""The report for one search, from the database alone (requirement 11).

Counts are queries over the outreach tables, never tallies carried in memory, so a
report built later matches one built at the time. It lists each Lead with its Decision
reasons and Messages, as Markdown, JSON or an Excel workbook. Emails and LinkedIn URLs
are masked, and a Lead's name is left out, unless ``reveal`` is asked for. The report
also says what was synthetic or offline (14.3).
"""

import io
import json
import uuid
from collections.abc import Iterable, Sequence
from datetime import date
from decimal import Decimal
from typing import Annotated

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import distinct, func, select
from sqlalchemy.orm import Session

from leadforge.lead_ingestion.store.lead_reader import load_lead
from leadforge.lead_ingestion.store.models import SourceRun
from leadforge.outreach.tables import (
    OutreachDecision,
    OutreachMessage,
    OutreachSearch,
    OutreachTriggerEvent,
)

__all__ = [
    "XLSX_TYPE",
    "EvidenceItem",
    "Funnel",
    "LeadRow",
    "Report",
    "build_report",
    "render_json",
    "render_markdown",
    "render_xlsx",
]

_MASK = "***"


class _Frozen(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid")


class Funnel(_Frozen):
    gathered: int
    selected: int
    invited: int
    emailed: int
    stalled: int
    rejected: int
    needs_enrichment: int
    manual_review: int


class EvidenceItem(_Frozen):
    """One cited finding. ``quote`` is untrusted page text: escape it wherever shown."""

    section: str  # "company_usage" or "person_fit"
    evidence_class: str
    source: str
    url: str
    observed_on: date | None
    quote: str
    relationship: str


class LeadRow(_Frozen):
    lead_id: uuid.UUID
    status: str
    score: Decimal
    reasons: tuple[str, ...]
    verdict: str | None = None
    company_usage: str | None = None
    company_usage_reason: str | None = None
    person_fit: str | None = None
    evidence: tuple[EvidenceItem, ...] = ()
    sequence: str
    name: str | None
    email: str | None
    linkedin_url: str | None
    invite: str | None
    email_subject: str | None
    email_body: str | None
    # The prompt or template version the stored Messages were written from. A run keeps
    # the wording it used, so this is what says a Message predates the current one.
    message_version: str | None = None


class Report(_Frozen):
    search_id: uuid.UUID
    mode: str
    query: str
    notes: tuple[str, ...]
    funnel: Funnel
    leads: Annotated[tuple[LeadRow, ...], Field(repr=False)]
    revealed: bool


def build_report(
    session: Session, search_id: uuid.UUID, *, reveal: bool = False
) -> Report:
    search = session.get_one(OutreachSearch, search_id)
    decisions = list(
        session.scalars(
            select(OutreachDecision)
            .where(OutreachDecision.search_id == search_id)
            .order_by(OutreachDecision.lead_id)
        )
    )
    return Report(
        search_id=search_id,
        mode=search.mode,
        query=search.query,
        notes=_notes(session, search),
        funnel=_funnel(session, search_id),
        leads=tuple(_row(session, d, reveal, search.mode) for d in decisions),
        revealed=reveal,
    )


def render_json(report: Report) -> str:
    return json.dumps(report.model_dump(mode="json"), indent=2, ensure_ascii=False)


def render_markdown(report: Report) -> str:
    f = report.funnel
    out = [
        f"# Outreach report: {report.mode}",
        "",
        f"Search `{report.search_id}`"
        + (" (details revealed)" if report.revealed else ""),
        "",
        "## What ran",
        *[f"- {n}" for n in report.notes],
        "",
        "## Funnel",
        "",
        "| Step | Count |",
        "|---|---|",
        *[f"| {k} | {v} |" for k, v in f.model_dump().items()],
        "",
        "## Leads",
    ]
    for row in report.leads:
        out += [
            "",
            f"### {row.lead_id} — {row.status} ({row.score})",
            f"- sequence: {row.sequence}",
            f"- email: {row.email or 'none'}; LinkedIn: {row.linkedin_url or 'none'}",
        ]
        if row.name:
            out.append(f"- name: {row.name}")
        out.append("- reasons: " + (", ".join(row.reasons) or "none"))
        if row.verdict is not None:
            out.append(f"- verdict: {row.verdict}")
        if row.company_usage is not None:
            why = f" ({row.company_usage_reason})" if row.company_usage_reason else ""
            out.append(f"- Company Usage: {row.company_usage}{why}")
        if row.person_fit is not None:
            out.append(f"- Person Fit: {row.person_fit}")
        out += _evidence_lines(row.evidence)
        if row.invite is not None:
            out += ["", "Invite:", "", _quote(row.invite)]
        if row.email_body is not None:
            out += [
                "",
                f"Email, subject: {row.email_subject}",
                "",
                _quote(row.email_body),
            ]
    return "\n".join(out) + "\n"


_LEAD_COLUMNS = (
    "Lead ID", "Name", "Email", "LinkedIn", "Status", "Score", "Verdict",
    "Company Usage", "Company Usage reason", "Person Fit", "Reasons", "Sequence",
    "Invite", "Email subject", "Email body",
)  # fmt: skip
_EVIDENCE_COLUMNS = (
    "Lead ID", "Section", "Class", "Relationship", "Observed on", "Source", "URL",
    "Quote",
)  # fmt: skip
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_STATUS_ORDER = {"selected": 0, "manual_review": 1}
_CELL_MAX = 32_767  # Excel's limit on one cell's text


def render_xlsx(report: Report) -> bytes:
    """The report as a workbook: Summary, Leads (selected first) and Evidence."""
    book = Workbook()
    summary = book.active
    assert summary is not None  # a new Workbook always has one sheet
    summary.title = "Summary"
    _fill(
        summary,
        [
            ("Mode", report.mode),
            ("Query", report.query),
            ("Search", str(report.search_id)),
            ("Details revealed", "yes" if report.revealed else "no"),
            (),
            ("Funnel step", "Count"),
            *report.funnel.model_dump().items(),
            (),
            ("What ran",),
            *[(n,) for n in report.notes],
        ],
    )
    leads = sorted(
        report.leads, key=lambda r: (_STATUS_ORDER.get(r.status, 2), -r.score)
    )
    _table(
        book.create_sheet("Leads"),
        _LEAD_COLUMNS,
        [
            (
                str(r.lead_id),
                r.name,
                r.email,
                r.linkedin_url,
                r.status,
                r.score,
                r.verdict,
                r.company_usage,
                r.company_usage_reason,
                r.person_fit,
                ", ".join(r.reasons),
                r.sequence,
                r.invite,
                r.email_subject,
                r.email_body,
            )
            for r in leads
        ],
    )
    _table(
        book.create_sheet("Evidence"),
        _EVIDENCE_COLUMNS,
        [
            (
                str(r.lead_id),
                i.section,
                i.evidence_class,
                i.relationship,
                i.observed_on,
                i.source,
                i.url,
                i.quote,
            )
            for r in leads
            for i in r.evidence
        ],
    )
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _table(
    sheet: Worksheet, header: Sequence[str], rows: Iterable[Sequence[object]]
) -> None:
    """A header row in bold, frozen and filterable, then the rows."""
    _fill(sheet, [header, *rows])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _fill(sheet: Worksheet, rows: Iterable[Sequence[object]]) -> None:
    """Write values only. Text never becomes a formula: quotes are untrusted page
    text, and a leading ``=`` would otherwise run when the file is opened."""
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = sheet.cell(r, c)
            if isinstance(value, str):
                cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)[:_CELL_MAX]
                cell.data_type = "s"
            else:
                cell.value = value  # type: ignore[assignment]


_SECTIONS = (("company_usage", "Company Usage"), ("person_fit", "Person Fit"))


def _evidence_lines(items: tuple[EvidenceItem, ...]) -> list[str]:
    out: list[str] = []
    for section, title in _SECTIONS:
        chosen = [i for i in items if i.section == section]
        if not chosen:
            continue
        out += ["", f"{title} evidence:"]
        for i in chosen:
            when = f", {i.observed_on}" if i.observed_on else ""
            link = (
                f"<{i.url}>"
                if i.url.lower().startswith(("http://", "https://"))
                else "no link"
            )
            out += [
                "",
                f"- {i.evidence_class} ({i.relationship}{when}) via {i.source}: {link}",
                "",
                _quote(i.quote),
            ]
    return out


def _quote(text: str) -> str:
    return "\n".join(f"> {line}" for line in text.splitlines())


def _funnel(session: Session, search_id: uuid.UUID) -> Funnel:
    decisions = select(OutreachDecision.id).where(
        OutreachDecision.search_id == search_id
    )

    def status(name: str) -> int:
        return (
            session.scalar(
                select(func.count())
                .select_from(OutreachDecision)
                .where(
                    OutreachDecision.search_id == search_id,
                    OutreachDecision.status == name,
                )
            )
            or 0
        )

    def events(*kinds: str) -> int:
        return (
            session.scalar(
                select(func.count(distinct(OutreachTriggerEvent.decision_id))).where(
                    OutreachTriggerEvent.decision_id.in_(decisions),
                    OutreachTriggerEvent.kind.in_(kinds),
                )
            )
            or 0
        )

    gathered = (
        session.scalar(
            select(func.count())
            .select_from(OutreachDecision)
            .where(OutreachDecision.search_id == search_id)
        )
        or 0
    )
    return Funnel(
        gathered=gathered,
        selected=status("selected"),
        invited=events("invite"),
        emailed=events("email", "fallback_email"),
        stalled=events("stalled"),
        rejected=status("rejected"),
        needs_enrichment=status("needs_enrichment"),
        manual_review=status("manual_review"),
    )


def _notes(session: Session, search: OutreachSearch) -> tuple[str, ...]:
    notes = [f"compiler: {search.compiler}"]
    if search.ingestion_run_id is not None:
        runs = session.execute(
            select(SourceRun.source_name, SourceRun.resolved_mode)
            .where(SourceRun.run_id == search.ingestion_run_id)
            .order_by(SourceRun.source_name)
        )
        notes += [f"source {name}: {mode}" for name, mode in runs]
    generators = sorted(
        set(
            session.scalars(
                select(OutreachMessage.generator)
                .join(
                    OutreachDecision, OutreachDecision.id == OutreachMessage.decision_id
                )
                .where(OutreachDecision.search_id == search.id)
            )
        )
    )
    notes.append("messages: " + (", ".join(generators) or "none written"))
    judged = session.scalar(
        select(func.count())
        .select_from(OutreachMessage)
        .join(OutreachDecision, OutreachDecision.id == OutreachMessage.decision_id)
        .where(
            OutreachDecision.search_id == search.id, OutreachMessage.judge.is_not(None)
        )
    )
    notes.append("judge: " + ("on" if judged else "off"))
    notes.append("dispatch: dry run, nothing is sent")
    return tuple(notes)


def _row(
    session: Session, decision: OutreachDecision, reveal: bool, mode: str
) -> LeadRow:
    messages = {
        m.variant: m
        for m in session.scalars(
            select(OutreachMessage).where(OutreachMessage.decision_id == decision.id)
        )
    }
    kinds = list(
        session.scalars(
            select(OutreachTriggerEvent.kind)
            .where(OutreachTriggerEvent.decision_id == decision.id)
            .order_by(OutreachTriggerEvent.at)
        )
    )
    stored = load_lead(session, decision.lead_id)
    email = (
        None if stored is None or stored.lead.email is None else str(stored.lead.email)
    )
    url = (
        None
        if stored is None or stored.lead.linkedin_url is None
        else str(stored.lead.linkedin_url)
    )
    invite, mail = messages.get("invite"), messages.get("email")
    return LeadRow(
        lead_id=decision.lead_id,
        status=decision.status,
        score=Decimal(decision.score_milli) / 1000,
        reasons=tuple(_reason(r) for r in decision.reasons),
        verdict=_verdict(decision.reasons) if mode == "users" else None,
        evidence=_evidence(decision.reasons) if mode == "users" else (),
        **(_grades(decision.reasons) if mode == "users" else {}),
        sequence=" > ".join(kinds) or "not started",
        name=stored.lead.full_name if reveal and stored is not None else None,
        email=email if reveal else _mask_email(email),
        linkedin_url=url if reveal else _mask_url(url),
        invite=None if invite is None else invite.body,
        email_subject=None if mail is None else mail.subject,
        email_body=None if mail is None else mail.body,
        message_version=next(
            (m.prompt_version for m in (invite, mail) if m is not None), None
        ),
    )


def _verdict(reasons: list[dict[str, object]]) -> str | None:
    """The verdict reason code: first in a users-mode Decision, ahead of score terms."""
    first = reasons[0] if reasons else {}
    if first.get("value") is None and first.get("weight") is None:
        return str(first["code"]) if "code" in first else None
    return None


def _grades(reasons: list[dict[str, object]]) -> dict[str, str | None]:
    """Company Usage and Person Fit grades from the verdict reason; None on old rows."""
    first = reasons[0] if reasons else {}
    keys = ("company_usage", "company_usage_reason", "person_fit")
    return {k: (None if first.get(k) is None else str(first[k])) for k in keys}


def _evidence(reasons: list[dict[str, object]]) -> tuple[EvidenceItem, ...]:
    items: list[EvidenceItem] = []
    for reason in reasons:
        for ref in reason.get("evidence_refs") or []:  # type: ignore[attr-defined]
            kind = str(ref["class"])
            observed = ref.get("observed_on")
            items.append(
                EvidenceItem(
                    section=(
                        "person_fit"
                        if kind == "person_self_stated"
                        else "company_usage"
                    ),
                    evidence_class=kind,
                    source=str(ref["source"]),
                    url=str(ref["url"]),
                    observed_on=None
                    if observed is None
                    else date.fromisoformat(observed),
                    quote=str(ref["quote"]),
                    relationship=str(ref["relationship"]),
                )
            )
    return tuple(items)


def _reason(reason: dict[str, object]) -> str:
    value = reason.get("value")
    if value is None:
        return str(reason["code"])
    try:
        shown = f"{float(str(value)):.2f}".rstrip("0").rstrip(".")
    except ValueError:
        shown = str(value)
    return f"{reason['code']}={shown}"


def _mask_email(address: str | None) -> str | None:
    if address is None:
        return None
    _, _, domain = address.partition("@")
    return f"{_MASK}@{domain}"


def _mask_url(url: str | None) -> str | None:
    if url is None:
        return None
    host = url.split("//", 1)[-1].split("/", 1)[0]
    return f"{host}/{_MASK}"
