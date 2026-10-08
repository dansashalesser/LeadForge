"""The outreach API and page, on the ingestion web app (13.1-13.5)."""

import io
import re
import shutil
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from leadforge.lead_ingestion.registry import SourceRegistry
from leadforge.lead_ingestion.web.api import create_app
from leadforge.outreach.report import XLSX_TYPE
from leadforge.outreach.web import create_router

CONFIG_DIR = Path(__file__).resolve().parents[4] / "config"
WRITE = {"X-LeadForge": "1"}


@pytest.fixture
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    shutil.copytree(CONFIG_DIR, tmp_path / "config")
    monkeypatch.chdir(tmp_path)
    discovered = SourceRegistry.discover()
    for source in discovered.names():
        for variable in discovered.source_class(source).required_env:
            monkeypatch.delenv(variable, raising=False)
    for name in ("LEADFORGE_MODE", "LEADFORGE_ENV_FILE", "LEADFORGE_MATCH_KEY_SECRET"):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    url = f"sqlite:///{tmp_path / 'store.db'}"
    monkeypatch.setenv("DATABASE_URL", url)
    app = create_app(
        main_url=url, demo_db=tmp_path / "demo.db", extra_routers=[create_router()]
    )
    return TestClient(app)


def _plan(
    client: TestClient, mode: str = "users", query: str = "mongodb"
) -> dict[str, object]:
    body: dict[str, object] = {"mode": mode, "query": query}
    if mode == "users":
        body["products"] = [query]
    answer = client.post("/api/outreach/plan", json=body, headers=WRITE)
    assert answer.status_code == 200, answer.text
    plan = answer.json()["plan"]
    assert isinstance(plan, dict)
    return plan


def _run(client: TestClient) -> str:
    started = client.post(
        "/api/outreach/searches", json={"plan": _plan(client)}, headers=WRITE
    )
    assert started.status_code == 200, started.text
    return str(started.json()["search_id"])


# Verifies: outreach requirements 13.4
@pytest.mark.parametrize(
    ("path", "body"),
    [
        ("/api/outreach/plan", {"mode": "users", "query": "x"}),
        ("/api/outreach/searches", {"plan": {}}),
        ("/api/outreach/tick", {}),
    ],
)
def test_every_write_is_refused_without_the_header(
    client: TestClient, path: str, body: dict[str, object]
) -> None:
    assert client.post(path, json=body).status_code == 403


# Verifies: outreach requirements 1.4
def test_a_plan_is_compiled_without_storing_or_spending(client: TestClient) -> None:
    plan = _plan(client, "free_text", "teams running mongodb")

    assert plan["compiler"] == "offline"
    assert plan["terms"] == ["mongodb"]
    assert client.get("/api/outreach/searches").status_code == 409  # no store yet


# Verifies: outreach requirements 1.1
# Verifies: outreach requirements 3.3
def test_plan_errors_are_named_and_returned_as_400(client: TestClient) -> None:
    bad_mode = client.post(
        "/api/outreach/plan", json={"mode": "everyone", "query": "x"}, headers=WRITE
    )
    no_domain = client.post(
        "/api/outreach/plan",
        json={"mode": "workers", "query": "Nowhere"},
        headers=WRITE,
    )

    assert bad_mode.status_code == no_domain.status_code == 400
    assert "unknown search mode" in bad_mode.json()["detail"]
    assert "no known domain" in no_domain.json()["detail"]


# Verifies: outreach requirements 13.4
def test_a_search_runs_lists_and_shows_detail_masked_unless_revealed(
    client: TestClient,
) -> None:
    search_id = _run(client)

    listed = client.get("/api/outreach/searches").json()["searches"]
    masked = client.get(f"/api/outreach/searches/{search_id}").json()["report"]
    shown = client.get(f"/api/outreach/searches/{search_id}?reveal=true").json()[
        "report"
    ]

    assert [s["search_id"] for s in listed] == [search_id]
    assert listed[0]["funnel"] == masked["funnel"]
    assert masked["leads"]
    assert all("***" in (r["email"] or "***") for r in masked["leads"])
    assert all(r["name"] is None for r in masked["leads"])
    assert any(r["name"] for r in shown["leads"])
    assert (
        client.get(
            f"/api/outreach/searches/{'0' * 8}-0000-0000-0000-{'0' * 12}"
        ).status_code
        == 404
    )


# Verifies: outreach requirements 13.5
# Verifies: outreach requirements 11.3
def test_the_markdown_download_matches_the_cli_report_and_the_json(
    client: TestClient,
) -> None:
    search_id = _run(client)

    md = client.get(f"/api/outreach/searches/{search_id}/report?format=md")
    js = client.get(f"/api/outreach/searches/{search_id}/report?format=json")

    assert md.headers["content-type"].startswith("text/markdown")
    assert js.headers["content-type"].startswith("application/json")
    for name, count in js.json()["funnel"].items():
        assert f"| {name} | {count} |" in md.text
    assert "## What ran" in md.text
    assert (
        client.get(f"/api/outreach/searches/{search_id}/report?format=pdf").status_code
        == 422
    )


def test_the_excel_download_has_every_lead_masked_unless_revealed(
    client: TestClient,
) -> None:
    search_id = _run(client)
    url = f"/api/outreach/searches/{search_id}/report"
    report = client.get(f"/api/outreach/searches/{search_id}").json()["report"]

    masked = client.get(url, params={"format": "xlsx"})
    revealed = client.get(url, params={"format": "xlsx", "reveal": True})

    assert masked.status_code == 200
    assert masked.headers["content-type"] == XLSX_TYPE
    assert "attachment" in masked.headers["content-disposition"]
    book = load_workbook(io.BytesIO(masked.content))
    assert book.sheetnames == ["Summary", "Leads", "Evidence"]
    rows = list(book["Leads"].iter_rows(values_only=True))
    assert rows[0][:3] == ("Lead ID", "Name", "Email")
    assert {r[0] for r in rows[1:]} == {lead["lead_id"] for lead in report["leads"]}
    emails = [r[2] for r in rows[1:] if r[2]]
    assert emails
    assert all(str(e).startswith("***@") for e in emails)
    assert all(r[1] is None for r in rows[1:])
    shown = list(load_workbook(io.BytesIO(revealed.content))["Leads"].values)
    assert not any("***" in str(r[2] or "") for r in shown[1:])


# Verifies: outreach requirements 13.3
def test_advancing_days_fires_what_became_due(client: TestClient) -> None:
    search_id = _run(client)

    early = client.post(
        "/api/outreach/tick",
        json={"search_id": search_id, "advance_days": 0},
        headers=WRITE,
    )
    late = client.post(
        "/api/outreach/tick",
        json={"search_id": search_id, "advance_days": 30},
        headers=WRITE,
    )
    both = client.post(
        "/api/outreach/tick",
        json={"advance_days": 1, "now": "2030-01-01T00:00:00Z"},
        headers=WRITE,
    )
    backwards = client.post(
        "/api/outreach/tick", json={"advance_days": -1}, headers=WRITE
    )

    assert early.status_code == late.status_code == 200
    assert early.json()["fired"] == 0
    assert both.status_code == backwards.status_code == 400


# Verifies: outreach requirements 13.1
# Verifies: outreach requirements 13.2
def test_the_page_has_the_search_panel_leads_timeline_and_download(
    client: TestClient,
) -> None:
    page = client.get("/")
    script = client.get("/outreach/app.js")
    styles = client.get("/outreach/app.css")

    assert page.status_code == script.status_code == styles.status_code == 200
    # The dashboard page carries the Search tab's script before its own, and its style.
    assert page.text.index("/outreach/app.js") < page.text.index("/static/app.js")
    assert "/outreach/app.css" in page.text
    for needle in ('"data-mode"', "Run search", "Review plan", "Advance", "download"):
        assert needle in script.text
    for mode in ("free_text", "workers", "users"):
        assert f'["{mode}"' in script.text
    assert script.text.index('["users"') < script.text.index('["workers"')
    assert 'key: "search"' in script.text
    assert "textContent" in script.text
    assert not re.search(r"innerHTML\s*=", script.text)
    assert "X-LeadForge" in script.text


# Verifies: outreach requirements 13.4
def test_the_ingestion_routes_still_work_beside_the_outreach_router(
    client: TestClient,
) -> None:
    assert client.get("/").status_code == 200
    assert client.get("/api/stores").status_code == 200


# Verifies: specs/user-recognition/requirements.md#2.2
def test_a_users_plan_with_no_product_is_a_400_naming_the_fault(
    client: TestClient,
) -> None:
    answer = client.post(
        "/api/outreach/plan", json={"mode": "users", "query": "x"}, headers=WRITE
    )

    assert answer.status_code == 400
    assert "catalog product" in answer.text


# Verifies: specs/user-recognition/requirements.md#2.3
def test_the_catalog_api_lists_every_filter_option(client: TestClient) -> None:
    options = client.get("/api/catalog")

    assert options.status_code == 200
    body = options.json()
    assert body["vendors"]
    for vendor in body["vendors"]:
        assert vendor["key"]
        assert vendor["name"]
        assert vendor["products"]
        assert {"key", "name"} <= set(vendor["products"][0])
        assert "ecosystem" in vendor
    assert body["role_families"]
    assert body["seniorities"]
    assert body["strictness"]["options"] == ["verified_only", "verified_plus_likely"]
    assert body["strictness"]["default"] in body["strictness"]["options"]
    assert body["max_evidence_age_days"] > 0
    assert set(body["budget"]) == {"searches", "fetches", "llm_calls"}


# Verifies: specs/user-recognition/requirements.md#2.4
def test_a_new_catalog_file_appears_in_the_api_with_no_frontend_change(
    client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("LEADFORGE_CATALOG_DIR", str(tmp_path / "config" / "catalog"))
    before = {v["key"] for v in client.get("/api/catalog").json()["vendors"]}
    (tmp_path / "config" / "catalog" / "newco.yaml").write_text(
        "vendor:\n  key: newco\n  name: NewCo\n  domains: [newco.example]\n"
        "products:\n  - key: widget\n    name: Widget\n"
        "    aliases:\n      - text: NewCo Widget\n",
        encoding="utf-8",
    )

    after = client.get("/api/catalog").json()["vendors"]

    assert "newco" not in before
    added = next(v for v in after if v["key"] == "newco")
    assert added["products"] == [{"key": "widget", "name": "Widget"}]


# Verifies: specs/user-recognition/requirements.md#2.4
def test_the_page_builds_its_filters_from_the_api(client: TestClient) -> None:
    script = client.get("/outreach/app.js").text

    assert '"aria-label": "Vendor"' in script
    assert '"aria-label": "Products"' in script
    assert "/api/catalog" in script
    assert "body.vendor" in script
    assert "products" in script
    assert "newco" not in script


# Verifies: specs/user-recognition/requirements.md#8.4
def test_the_page_renders_evidence_with_escaped_text_and_safe_links(
    client: TestClient,
) -> None:
    script = client.get("/outreach/app.js").text

    assert "evidence" in script
    assert "noopener noreferrer" in script
    assert "https?:" in script
    assert not re.search(r"innerHTML\s*=", script)


def _draft_file(tmp_path: Path, key: str = "newco") -> Path:
    path = tmp_path / "config" / "catalog" / "drafts" / f"{key}.yaml"
    path.parent.mkdir(exist_ok=True)
    path.write_text(
        f"vendor: {{key: {key}, name: NewCo}}\n"
        "products:\n  - key: widget\n    aliases: [{text: NewCo Widget}]\n",
        encoding="utf-8",
    )
    return path


# Verifies: specs/user-recognition/requirements.md#2.5
def test_a_draft_is_refused_by_a_search_until_approved(
    client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("LEADFORGE_CATALOG_DIR", str(tmp_path / "config" / "catalog"))
    _draft_file(tmp_path)
    body = {
        "mode": "users",
        "query": "newco",
        "vendor": "newco",
        "products": ["widget"],
    }

    refused = client.post("/api/outreach/plan", json=body, headers=WRITE)
    assert refused.status_code == 400
    assert "unapproved draft" in refused.text
    assert client.post("/api/catalog/drafts/newco/approve").status_code in (401, 403)

    approved = client.post("/api/catalog/drafts/newco/approve", headers=WRITE)
    assert approved.status_code == 200, approved.text
    assert (
        client.post("/api/outreach/plan", json=body, headers=WRITE).status_code == 200
    )
    assert (
        client.post("/api/catalog/drafts/newco/approve", headers=WRITE).status_code
        == 404
    )


# The dashboard's scripts are versioned by file content, so a browser that cached an
# older app.js (the server sends no cache headers) cannot run it without the Search tab.
def test_the_dashboard_versions_its_scripts_and_is_never_cached(
    client: TestClient,
) -> None:
    response = client.get("/")

    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    assert '<script src="/outreach/app.js?v=' in response.text
    assert '<script src="/static/app.js?v=' in response.text
    assert '<link rel="stylesheet" href="/static/app.css?v=' in response.text
    assert '<link rel="stylesheet" href="/outreach/app.css?v=' in response.text
    # The version changes when the file does, so a new build is fetched.
    assert client.get("/").text == response.text


def test_a_lead_s_messages_are_written_on_demand_and_not_stored(
    client: TestClient,
) -> None:
    search_id = _run(client)
    report = client.get(f"/api/outreach/searches/{search_id}").json()["report"]
    lead_id = report["leads"][0]["lead_id"]

    answer = client.post(
        f"/api/outreach/leads/{lead_id}/messages", json={}, headers=WRITE
    )
    again = client.get(f"/api/outreach/searches/{search_id}").json()["report"]

    assert answer.status_code == 200, answer.text
    body = answer.json()
    assert body["generator"] == "offline"
    assert {d["kind"] for d in body["drafts"]} == {"invite", "email"}
    assert all(d["body"] for d in body["drafts"])
    assert again["leads"] == report["leads"]


def test_on_demand_messages_need_the_write_header_and_a_known_lead(
    client: TestClient,
) -> None:
    _run(client)
    missing = "00000000-0000-0000-0000-000000000000"

    assert (
        client.post(f"/api/outreach/leads/{missing}/messages", json={}).status_code
        == 403
    )
    unknown = client.post(
        f"/api/outreach/leads/{missing}/messages", json={}, headers=WRITE
    )
    assert unknown.status_code == 404
